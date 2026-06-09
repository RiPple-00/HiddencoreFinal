"""Korean Excel workbooks for labeling and project reference."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.datavalidation import DataValidation

import schema as S
import settings as CFG
from labeling.gt_io import load_all_ground_truths
from preprocessing.gallery_index import image_id_to_index, load_gallery_index

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF")
FREEZE_ROW = 2

MANUAL_LABEL_CHOICES = (
    "patient_1",
    "patient_2",
    "patient_3",
    "patient_4",
    "patient_5",
    "patient_6",
    "unknown",
    "no_face",
    "ignore",
)

LABELING_COLUMNS_KO = [
    ("global_index", "전역번호"),
    ("image_id", "이미지ID"),
    ("program_label", "프로그램"),
    ("program_id", "프로그램ID"),
    ("face_index", "얼굴번호"),
    ("status", "예측상태"),
    ("predicted_patient", "예측입소자"),
    ("top1_similarity", "1등유사도"),
    ("top2_patient", "2등입소자"),
    ("top2_similarity", "2등유사도"),
    ("similarity_gap", "점수차"),
    ("patient_1_similarity", "P1유사도"),
    ("patient_2_similarity", "P2유사도"),
    ("patient_3_similarity", "P3유사도"),
    ("patient_4_similarity", "P4유사도"),
    ("patient_5_similarity", "P5유사도"),
    ("patient_6_similarity", "P6유사도"),
    ("gt_label", "정답_GT"),
    ("manual_label", "수동라벨_입력"),
    ("label_status", "라벨상태"),
    ("review_flag", "검수필요"),
    ("visualized_file", "시각화파일"),
    ("aligned_face_path", "얼굴crop경로"),
    ("memo", "메모"),
]

IMAGE_SUMMARY_KO = [
    ("global_index", "전역번호"),
    ("image_id", "이미지ID"),
    ("program_label", "프로그램"),
    ("face_count", "얼굴수"),
    ("matched_count", "MATCHED수"),
    ("unknown_count", "UNKNOWN수"),
    ("uncertain_count", "UNCERTAIN수"),
    ("low_conf_count", "LOW_CONF수"),
    ("no_face_count", "NO_FACE수"),
    ("failed_count", "FAILED수"),
    ("top_predicted", "대표예측"),
    ("visualized_file", "시각화파일"),
    ("gt_done", "GT완료여부"),
]


def _gt_map() -> dict[tuple[str, int], str | None]:
    """(image_id, person_index) -> patient_id_gt"""
    idx_to_ids: dict[int, set[str]] = {}
    for e in load_gallery_index().get("entries", []):
        idx = int(e["index"])
        ids = {str(e.get("image_id", "")).replace("\\", "/")}
        for key in ("legacy_image_id", "original_image_id", "gallery_path", "preprocessed_path"):
            v = e.get(key)
            if v:
                ids.add(str(v).replace("\\", "/"))
        idx_to_ids[idx] = ids

    out: dict[tuple[str, int], str | None] = {}
    for gt in load_all_ground_truths():
        img = str(gt.get("image_id", "")).replace("\\", "/")
        status = gt.get("label_status")
        for p in gt.get("persons", []):
            pi = int(p.get("person_index", 0))
            label = p.get("patient_id_gt")
            keys = {(img, pi)}
            for idx, id_set in idx_to_ids.items():
                if img in id_set:
                    for alt in id_set:
                        keys.add((alt, pi))
            for k in keys:
                out[k] = label
    return out


def _gt_status_map() -> dict[str, str]:
    m: dict[str, str] = {}
    for gt in load_all_ground_truths():
        img = str(gt.get("image_id", "")).replace("\\", "/")
        m[img] = gt.get("label_status", "pending")
    return m


def enrich_for_labeling(df: pd.DataFrame) -> pd.DataFrame:
    """Add global_index, GT, visualization paths for labeling sheets."""
    id_map = image_id_to_index()
    gt = _gt_map()
    gt_status = _gt_status_map()

    out = df.copy()
    out["image_id_norm"] = out["image_id"].astype(str).str.replace("\\", "/", regex=False)
    out["global_index"] = out["image_id_norm"].map(id_map)

    def lookup_gt(row) -> str | None:
        img = row["image_id_norm"]
        fi = int(row["face_index"]) if pd.notna(row["face_index"]) else 0
        return gt.get((img, fi))

    out["gt_label"] = out.apply(lookup_gt, axis=1)
    out["label_status"] = out["image_id_norm"].map(gt_status).fillna("pending")
    out["visualized_file"] = out["global_index"].apply(
        lambda x: f"output/visualized/gallery_preprocessed/{int(x)}.png"
        if pd.notna(x)
        else ""
    )

    sort_cols = ["global_index", "face_index"]
    out = out.sort_values(
        [c for c in sort_cols if c in out.columns],
        na_position="last",
    )
    return out


def _rename_columns(df: pd.DataFrame, mapping: list[tuple[str, str]]) -> pd.DataFrame:
    cols = [src for src, _ in mapping if src in df.columns]
    sub = df[cols].copy()
    sub.columns = [dst for src, dst in mapping if src in df.columns]
    return sub


def _write_rows_sheet(ws, rows: list[tuple], max_width_rows: int = 300) -> None:
    for r_idx, row in enumerate(rows, 1):
        for c_idx, val in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=val)
    n_cols = len(rows[0]) if rows else 1
    _style_sheet(ws, n_cols, max_width_rows=max_width_rows)


def _style_sheet(ws, n_cols: int, validation_col: str | None = None, max_width_rows: int = 200) -> None:
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = f"A{FREEZE_ROW}"
    ws.auto_filter.ref = ws.dimensions

    for col in range(1, n_cols + 1):
        letter = get_column_letter(col)
        max_len = 12
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_width_rows), min_col=col, max_col=col):
            val = row[0].value
            if val is not None:
                max_len = max(max_len, min(len(str(val)) + 2, 48))
        ws.column_dimensions[letter].width = max_len

    if validation_col:
        col_idx = None
        for c in range(1, n_cols + 1):
            if ws.cell(row=1, column=c).value == validation_col:
                col_idx = c
                break
        if col_idx:
            letter = get_column_letter(col_idx)
            dv = DataValidation(
                type="list",
                formula1=f'"{",".join(MANUAL_LABEL_CHOICES)}"',
                allow_blank=True,
            )
            dv.error = "patient_1~6, unknown, no_face, ignore 중 선택"
            dv.errorTitle = "수동라벨"
            last = max(ws.max_row, 2)
            dv.add(f"{letter}2:{letter}{last}")
            ws.add_data_validation(dv)


def build_image_summary(df: pd.DataFrame) -> pd.DataFrame:
    enriched = enrich_for_labeling(df)
    rows: list[dict] = []
    for image_id, grp in enriched.groupby("image_id_norm", sort=False):
        g = grp.iloc[0]
        status_counts = grp["status"].value_counts().to_dict()
        preds = grp[grp["status"] == S.IDENTITY_MATCHED]["predicted_patient"]
        rows.append(
            {
                "global_index": g.get("global_index"),
                "image_id": image_id,
                "program_label": g.get("program_label"),
                "face_count": len(grp),
                "matched_count": status_counts.get(S.IDENTITY_MATCHED, 0),
                "unknown_count": status_counts.get(S.IDENTITY_UNKNOWN, 0),
                "uncertain_count": status_counts.get(S.IDENTITY_UNCERTAIN, 0),
                "low_conf_count": status_counts.get(S.IDENTITY_LOW_CONFIDENCE, 0),
                "no_face_count": status_counts.get(S.IDENTITY_NO_FACE, 0),
                "failed_count": status_counts.get(S.IDENTITY_FAILED, 0),
                "top_predicted": preds.iloc[0] if len(preds) else grp.iloc[0].get("predicted_patient"),
                "visualized_file": g.get("visualized_file"),
                "gt_done": g.get("label_status") == "done",
            }
        )
    summary = pd.DataFrame(rows)
    if "global_index" in summary.columns:
        summary = summary.sort_values("global_index", na_position="last")
    return summary


def _safe_save_workbook(wb, path: Path) -> Path:
    try:
        wb.save(path)
        return path
    except PermissionError:
        alt = path.with_name(f"{path.stem}_new{path.suffix}")
        wb.save(alt)
        return alt


def export_labeling_workbook(df: pd.DataFrame, path: Path | None = None) -> Path:
    """
    Main labeling Excel:
      - 라벨링_작성: face-level, Korean headers, dropdown on 수동라벨
      - 이미지별요약
      - 검수필요: uncertain/unknown/low_confidence
      - 전체데이터
    """
    path = path or CFG.OUTPUT_REPORTS_DIR / "유사도_라벨링.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)

    enriched = enrich_for_labeling(df)
    if "manual_label" not in enriched.columns:
        enriched["manual_label"] = None
    if "memo" not in enriched.columns:
        enriched["memo"] = None

    review = enriched[
        enriched["status"].isin(
            [
                S.IDENTITY_UNCERTAIN,
                S.IDENTITY_UNKNOWN,
                S.IDENTITY_LOW_CONFIDENCE,
            ]
        )
        | (enriched["review_flag"] == True)  # noqa: E712
    ].copy()

    summary = build_image_summary(df)

    wb = Workbook()
    wb.remove(wb.active)

    ws1 = wb.create_sheet("라벨링_작성", 0)
    labeling_df = _rename_columns(enriched, LABELING_COLUMNS_KO)
    for r_idx, row in enumerate(dataframe_to_rows(labeling_df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            ws1.cell(row=r_idx, column=c_idx, value=value)
    _style_sheet(ws1, len(labeling_df.columns), validation_col="수동라벨_입력")

    ws2 = wb.create_sheet("이미지별요약", 1)
    summary_df = _rename_columns(summary, IMAGE_SUMMARY_KO)
    for r_idx, row in enumerate(dataframe_to_rows(summary_df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            ws2.cell(row=r_idx, column=c_idx, value=value)
    _style_sheet(ws2, len(summary_df.columns))

    ws3 = wb.create_sheet("검수필요", 2)
    review_df = _rename_columns(review, LABELING_COLUMNS_KO)
    for r_idx, row in enumerate(dataframe_to_rows(review_df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            ws3.cell(row=r_idx, column=c_idx, value=value)
    _style_sheet(ws3, len(review_df.columns), validation_col="수동라벨_입력")

    ws4 = wb.create_sheet("전체데이터", 3)
    full_ko = LABELING_COLUMNS_KO + [
        ("source_image", "원본경로"),
        ("det_score", "검출점수"),
        ("top1_patient", "1등입소자"),
        ("patient_similarities", "유사도JSON"),
    ]
    full_df = _rename_columns(enriched, full_ko)
    for r_idx, row in enumerate(dataframe_to_rows(full_df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            ws4.cell(row=r_idx, column=c_idx, value=value)
    _style_sheet(ws4, len(full_df.columns))

    # Status sheets
    for status, title in [
        (S.IDENTITY_MATCHED, "MATCHED"),
        (S.IDENTITY_UNKNOWN, "UNKNOWN"),
        (S.IDENTITY_UNCERTAIN, "UNCERTAIN"),
        (S.IDENTITY_LOW_CONFIDENCE, "LOW_CONF"),
    ]:
        sub = enriched[enriched["status"] == status]
        if len(sub) == 0:
            continue
        ws = wb.create_sheet(title[:31])
        sub_df = _rename_columns(sub, LABELING_COLUMNS_KO)
        for r_idx, row in enumerate(dataframe_to_rows(sub_df, index=False, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                ws.cell(row=r_idx, column=c_idx, value=value)
        _style_sheet(ws, len(sub_df.columns))

    return _safe_save_workbook(wb, path)


def export_project_guide_workbook(path: Path | None = None) -> Path:
    """Full project file catalog for non-developers (every file under project root)."""
    from face.project_file_catalog import entries_to_rows, scan_project

    path = path or CFG.OUTPUT_REPORTS_DIR / "프로젝트_파일_가이드.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)

    entries = scan_project(CFG.ROOT)
    all_rows = entries_to_rows(entries)

    wb = Workbook()
    wb.remove(wb.active)

    ws_all = wb.create_sheet("전체파일목록", 0)
    _write_rows_sheet(ws_all, all_rows, max_width_rows=400)

    def _filter_sheet(title: str, pred, idx: int) -> None:
        sub = [all_rows[0]] + [
            r for r in all_rows[1:] if pred(r[0], r[1])
        ]
        if len(sub) <= 1:
            return
        ws = wb.create_sheet(title[:31], idx)
        _write_rows_sheet(ws, sub, max_width_rows=200)

    _filter_sheet(
        "코드파일_py",
        lambda p, c: p.endswith(".py") and "/__pycache__/" not in p,
        1,
    )
    _filter_sheet(
        "Python캐시_pyc",
        lambda p, c: c == "Python캐시" or "/__pycache__/" in p,
        2,
    )
    _filter_sheet(
        "data폴더",
        lambda p, c: p.startswith("data/"),
        3,
    )
    _filter_sheet(
        "output산출물",
        lambda p, c: p.startswith("output/"),
        4,
    )
    _filter_sheet(
        "문서_스키마",
        lambda p, c: p.startswith("docs/") or p.startswith("schemas/") or p.endswith(".md"),
        5,
    )

    # Folder summary
    from collections import defaultdict

    folder_stats: dict[str, dict] = defaultdict(lambda: {"count": 0, "categories": set(), "preserve": set()})
    for e in entries:
        parts = e.rel_path.split("/")
        key = parts[0] if len(parts) == 1 else "/".join(parts[:2])
        folder_stats[key]["count"] += 1
        folder_stats[key]["categories"].add(e.category)
        folder_stats[key]["preserve"].add(e.preserve)

    ws_fold = wb.create_sheet("폴더요약", 6)
    fold_rows = [
        ("폴더(상위2단계)", "파일수", "포함분류", "보존특성"),
    ]
    for key in sorted(folder_stats.keys(), key=lambda x: (x.count("/"), x)):
        st = folder_stats[key]
        fold_rows.append(
            (
                key,
                st["count"],
                ", ".join(sorted(st["categories"])),
                ", ".join(sorted(st["preserve"])),
            )
        )
    _write_rows_sheet(ws_fold, fold_rows)

    ws2 = wb.create_sheet("파이프라인순서", 7)
    pipeline = [
        ("순서", "명령", "입력", "출력"),
        (1, "python main.py --step preprocess", "gallery, patients", "*_preprocessed"),
        (2, "python main.py --step build-patient-db", "patients_preprocessed", "face_db"),
        (3, "python main.py --step detect-faces", "gallery_preprocessed", "activity_faces"),
        (4, "python main.py --step match", "activity_faces + face_db", "reports, visualized"),
        (5, "scripts/export_excel_reports.py", "face_similarity_results.csv", "유사도_라벨링.xlsx"),
        (6, "라벨링 후", "유사도_라벨링.xlsx 수동라벨", "ground_truth JSON"),
        (7, "python scripts/eval_identity.py", "ground_truth + 예측", "eval 리포트"),
    ]
    for r_idx, row in enumerate(pipeline, 1):
        for c_idx, val in enumerate(row, 1):
            ws2.cell(row=r_idx, column=c_idx, value=val)
    _style_sheet(ws2, 4)

    ws3 = wb.create_sheet("라벨값안내", 8)
    labels = [
        ("수동라벨_입력", "의미"),
        ("patient_1 ~ patient_6", "해당 등록 입소자"),
        ("unknown", "얼굴은 있으나 입소자 특정 불가"),
        ("no_face", "이 얼굴/사람 없음"),
        ("ignore", "평가에서 제외"),
        ("예측상태 MATCHED", "threshold 통과, 1등 확정"),
        ("예측상태 UNCERTAIN", "1·2등 점수 차이 작음"),
        ("예측상태 UNKNOWN", "threshold 미달"),
        ("예측상태 LOW_CONFIDENCE", "낮은 신뢰 구간"),
    ]
    for r_idx, row in enumerate(labels, 1):
        for c_idx, val in enumerate(row, 1):
            ws3.cell(row=r_idx, column=c_idx, value=val)
    _style_sheet(ws3, 2)

    ws4 = wb.create_sheet("읽는법", 9)
    guide = [
        ("항목", "설명"),
        ("전체파일목록", "프로젝트 루트 아래 모든 파일(이미지·캐시·코드 포함). Ctrl+F로 경로 검색"),
        ("코드파일_py", ".py 소스만. 내용요약에 class/def 목록 포함"),
        ("Python캐시_pyc", "__pycache__/*.pyc — 대응 .py의 자동 생성 바이트코드. 삭제 가능"),
        ("data폴더", "입력 데이터·중간 산출물 (gallery, face_db, activity_faces 등)"),
        ("output산출물", "match/export 후 CSV·PNG·Excel·로그"),
        ("보존/삭제", "보존=직접 관리 | 재생성=main.py로 다시 생성 | 삭제가능=캐시"),
        ("폴더요약", "상위 2단계 폴더별 파일 개수"),
    ]
    for r_idx, row in enumerate(guide, 1):
        for c_idx, val in enumerate(row, 1):
            ws4.cell(row=r_idx, column=c_idx, value=val)
    _style_sheet(ws4, 2)

    return _safe_save_workbook(wb, path)


def export_all_excel_reports(df: pd.DataFrame) -> dict[str, Path]:
    """Labeling workbook + project guide."""
    paths = {
        "유사도_라벨링.xlsx": export_labeling_workbook(df),
        "프로젝트_파일_가이드.xlsx": export_project_guide_workbook(),
    }
    return paths
